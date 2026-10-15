"""Generate an academically integrated Chapter 2 Word document.

This script uses the final focused experiments as the main paper evidence.
Representative algorithms and stress-test figures are integrated into the
original 2.6 and 2.7 structure.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt
except ModuleNotFoundError as exc:
    raise ModuleNotFoundError(
        "Generating the Word chapter requires python-docx. "
        "Install it in the MMGN conda environment with: "
        "conda install -n mmgn -c conda-forge python-docx"
    ) from exc

REPO_ROOT = Path(__file__).resolve().parents[4]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from experiments.chapter2.noiseEnhance.data_utils import load_flow_npz
from experiments.chapter2.noiseEnhance.methods import METHOD_LABELS, METHODS


ROOT = Path(__file__).resolve().parents[1]
OUT = REPO_ROOT / "outputs" / "chapter2"
FIG = OUT / "figures"
METRIC = OUT / "metrics"
DATA_PATH = REPO_ROOT / "data" / "current" / "bohai.npz"
DOCX = ROOT / "第二章_含噪稀疏观测建模与噪声增强算法_论文整合版.docx"


METHOD_ORDER = list(METHODS)
METHOD_NAMES = METHOD_LABELS


def set_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.text = ""
    run = p.add_run(text)
    set_font(run, bold=bold)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.append(borders)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.bold = True
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(14)
    doc.styles["Heading 3"].font.size = Pt(12)
    return doc


def add_paragraph(doc: Document, text: str) -> None:
    for block in text.strip().split("\n"):
        block = block.strip()
        if not block:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(block)
        set_font(run)


def add_equation(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_font(run, size=10.5)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_font(run, bold=True)


def add_picture(doc: Document, filename: str, caption: str, note: str, width: float = 5.85) -> None:
    path = FIG / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)
    add_paragraph(doc, note)


def add_table(doc: Document, title: str, headers: list[str], rows: list[list[str]]) -> None:
    add_caption(doc, title)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    set_table_borders(table)
    doc.add_paragraph()


def fmt(value: float, digits: int = 4) -> str:
    return f"{float(value):.{digits}f}"


def pct(value: float) -> str:
    return f"{float(value):.2f}"


def load_metrics() -> dict[str, pd.DataFrame]:
    required = {
        "ranking": "focused_ranking.csv",
        "noise": "focused_noise_ratio.csv",
        "sampling": "focused_sampling_rate.csv",
        "outlier": "focused_outlier_fraction.csv",
        "noise_model": "focused_noise_model.csv",
        "parameter": "focused_parameter_sensitivity.csv",
        "gap": "focused_gap_table.csv",
    }
    result = {}
    for key, filename in required.items():
        path = METRIC / filename
        if not path.exists():
            raise FileNotFoundError(f"缺少实验结果文件：{path}")
        result[key] = pd.read_csv(path)
    return result


def rows_ranking(ranking: pd.DataFrame) -> list[list[str]]:
    amr_rmse = float(ranking.loc[ranking["method"] == "amr_gtv", "rmse_mean"].iloc[0])
    rows = []
    for _, row in ranking.iterrows():
        reduction = (float(row["rmse_mean"]) - amr_rmse) / float(row["rmse_mean"]) * 100.0
        rows.append(
            [
                METHOD_NAMES.get(row["method"], row["method_label"]),
                fmt(row["rmse_mean"]),
                fmt(row["rmse_std"]),
                pct(row["gain_mean"]),
                pct(reduction),
            ]
        )
    return rows


def pivot_rows(df: pd.DataFrame, index_col: str, value_col: str) -> list[list[str]]:
    pivot = df.pivot(index=index_col, columns="method", values=value_col)
    rows = []
    for idx in pivot.index:
        row = [f"{idx:g}" if isinstance(idx, (float, np.floating)) else str(idx)]
        for method in METHOD_ORDER:
            row.append(fmt(pivot.loc[idx, method]) if method in pivot.columns else "-")
        rows.append(row)
    return rows


def build_doc() -> Path:
    metrics = load_metrics()
    data = load_flow_npz(DATA_PATH)
    valid_points = int(np.sum(data.mask[0]))
    total_points = int(data.field.shape[1] * data.field.shape[2])
    ranking = metrics["ranking"].sort_values("rmse_mean")
    amr = ranking[ranking["method"] == "amr_gtv"].iloc[0]
    median = ranking[ranking["method"] == "median"].iloc[0]
    graph = ranking[ranking["method"] == "graph_laplacian"].iloc[0]
    ragle = ranking[ranking["method"] == "ragle"].iloc[0]
    noisy = ranking[ranking["method"] == "noisy"].iloc[0]

    doc = setup_document()
    doc.add_heading("第2章 含噪稀疏观测建模与噪声增强算法", level=1)

    doc.add_heading("2.1 引言", level=2)
    add_paragraph(
        doc,
        """
        第一章已经指出，实际海洋流场观测往往同时具有空间采样稀疏、观测误差不均匀和局部异常值等特点。若直接将此类观测输入后续连续场重建模型，模型会把局部噪声视为真实物理结构，从而降低重建精度并影响未来流场预测。因此，在重建模型之前对稀疏观测点进行噪声建模与增强，是提高整体流场建模可靠性的关键环节。
        本章研究对象不是完整规则网格上的图像去噪，而是不规则稀疏观测点集合上的数值修正。为适应该问题，本章首先建立含噪稀疏观测的数据表示和噪声模型，然后利用k近邻图刻画观测点之间的局部空间关系。在此基础上，本文提出自适应多尺度鲁棒图总变差增强算法AMR-GTV。该方法不假设观测者已知各点真实噪声强度，而是根据含噪观测本身估计点级置信度，并结合多尺度图约束、自适应中值先验和鲁棒边缘保持项完成观测增强。
        本章实验部分选取含噪观测、KNN均值、中值滤波、图拉普拉斯、RAGLE基础版和AMR-GTV六种具有代表性的方法构成主对比体系。该选择覆盖了无处理基线、传统局部统计方法、经典图平滑方法、鲁棒图拉普拉斯方法和本文改进方法，能够较清晰地反映不同建模思想在困难观测场景下的差异。
        """,
    )

    doc.add_heading("2.2 稀疏观测流场数据表示", level=2)
    doc.add_heading("2.2.1 连续流场与离散观测点表示", level=3)
    add_paragraph(
        doc,
        f"""
        设研究区域内某一时刻的连续流场标量物理量为u(x)，其中x表示二维空间坐标。实际数据以规则网格形式存储，但本章只在有效海域掩膜内抽取少量观测点，形成稀疏观测集合O_t={{(x_i,y_i)}}。本实验数据的单帧网格规模为{data.field.shape[1]}×{data.field.shape[2]}，有效海域点数为{valid_points}，占全网格点数{total_points}的一部分。后续增强算法仅在稀疏点集合上运行，不使用未采样网格点的真实物理量。
        """,
    )
    add_equation(doc, "O_t = {(x_i, y_i) | i = 1, 2, ..., N, x_i ∈ Ω_valid}")
    add_picture(
        doc,
        "fig_2_1_dataset_mask_sampling.png",
        "图2-1 数据场、海域掩膜与稀疏采样点示意",
        "图2-1给出了完整流场、有效海域掩膜和稀疏采样点之间的关系。该图说明，本章增强算法面对的是空间上不规则分布的有限观测点，而不是完整网格图像。因此，算法既需要利用局部空间连续性，又必须避免依赖规则卷积或全场真实值。",
    )

    doc.add_heading("2.2.2 稀疏采样方式定义", level=3)
    add_paragraph(
        doc,
        """
        为模拟不同观测密度，本章按照采样率从有效海域掩膜内随机抽取观测点。若采样率为r，单帧网格总点数为H×W，则观测点数量近似为N=round(rHW)，并限制所有观测点均位于海域掩膜内部。实验采用多个随机种子重复采样，以降低单次采样点分布对结果评价的偶然影响。
        """,
    )
    doc.add_heading("2.2.3 观测点坐标与物理量归一化处理", level=3)
    add_paragraph(
        doc,
        """
        为避免经纬度量纲差异影响距离权重计算，本文将纬度和经度分别归一化至[0,1]区间，并以归一化坐标构建k近邻图。物理量采用数据文件中的temperature字段表示。所有误差指标均在增强后的稀疏观测点上计算，即比较增强值与采样点处干净观测值之间的差异。
        """,
    )

    doc.add_heading("2.3 含噪稀疏观测问题建模", level=2)
    doc.add_heading("2.3.1 加性噪声模型", level=3)
    add_paragraph(
        doc,
        """
        设干净观测值为y_i^*，含噪观测值为y_i。本文采用加性噪声模型描述观测误差，即噪声直接作用于物理量数值而不改变观测点坐标。与理想同方差高斯噪声不同，真实传感器误差常表现为不同位置噪声强度不一致，并可能出现少量离群观测点。因此，本章实验同时考虑同方差高斯噪声、异方差高斯噪声和异方差离群噪声。
        """,
    )
    add_equation(doc, "y_i = y_i^* + ε_i,     ε_i ~ N(0, σ_i^2)")
    add_picture(
        doc,
        "fig_2_2_noise_model.png",
        "图2-2 含噪观测模型及噪声残差示意",
        "图2-2展示干净观测、含噪观测和噪声残差的空间分布。可以看到，噪声残差并非总是均匀分布，局部区域可能出现明显扰动。该现象说明，对所有观测点使用相同保真权重并不合理，算法需要区分相对可信观测点与疑似异常观测点。",
    )
    doc.add_heading("2.3.2 不同噪声强度设置", level=3)
    add_paragraph(
        doc,
        """
        主压力实验采用噪声强度0.40、0.60、0.80和1.00，并固定离群点比例为0.10，以构造能够区分不同算法鲁棒性的困难观测场景。低噪声条件下，多数方法的差距通常较小，甚至普通平滑可能因过度处理而造成误差增加；中高噪声和离群点条件更能体现鲁棒图增强方法的实际价值。
        """,
    )
    add_picture(
        doc,
        "fig_2_3_noise_model_comparison.png",
        "图2-3 不同未知噪声模型下的代表算法对比",
        "图2-3比较了同方差高斯、异方差高斯和异方差离群噪声三种模型。结果表明，在单纯高斯噪声下，RAGLE基础版等图方法具有竞争力；而当噪声中包含离群点时，鲁棒局部统计和AMR-GTV更稳定。这也说明本文方法主要面向复杂观测误差，并不主张在所有轻噪声场景中强行平滑。",
    )
    doc.add_heading("2.3.3 含噪观测对流场重建的影响分析", level=3)
    add_paragraph(
        doc,
        """
        含噪观测对后续重建模型的影响主要体现在两个方面。第一，局部大误差观测会破坏稀疏点集合的空间一致性，使隐式神经表示在拟合时产生局部伪结构。第二，异方差噪声会使不同观测点可靠性不同，若训练时仍赋予相同权重，高噪声点会对模型参数产生不成比例的影响。因此，本章增强算法的目标是在保持原始观测空间分布的前提下，降低数值噪声并弱化异常观测点影响。
        """,
    )

    doc.add_heading("2.4 基于邻域图的观测点关系建模", level=2)
    doc.add_heading("2.4.1 k-近邻图构建方法", level=3)
    add_paragraph(
        doc,
        """
        稀疏观测点不具备规则网格拓扑，因此本文通过k近邻图建立点间关系。对每个观测点x_i，选取欧氏距离最近的k个观测点作为邻居，构成边集合E。该图结构提供了在不规则点集上定义平滑、边权和图拉普拉斯矩阵的基础。
        """,
    )
    add_picture(
        doc,
        "fig_2_4_knn_graph.png",
        "图2-4 稀疏观测点的局部kNN图结构",
        "图2-4显示了观测点之间的局部kNN连接关系。空间邻近点通常具有相似物理量，但在边界区域或异常噪声点附近也可能出现显著差异。因此，后续算法不能只依赖空间距离，还需要根据观测值差异和点级置信度调整平滑强度。",
    )
    doc.add_heading("2.4.2 空间距离权重定义", level=3)
    add_paragraph(
        doc,
        """
        对图中相邻节点(i,j)，本文采用高斯核定义基础空间权重，其中d_ij为归一化坐标空间中的欧氏距离，τ为由邻域距离中位数估计的尺度参数。距离越近，空间权重越大；距离越远，空间权重越小。
        """,
    )
    add_equation(doc, "w_ij = exp(-d_ij^2 / (2τ^2)),     (i,j) ∈ E")
    doc.add_heading("2.4.3 图邻接矩阵与图拉普拉斯矩阵构造", level=3)
    add_paragraph(
        doc,
        """
        由边权w_ij可构造图邻接矩阵W，并进一步得到图拉普拉斯矩阵L=D-W，其中D为节点度矩阵。图拉普拉斯二次型z^TLz可度量图信号z在相邻节点之间的整体不平滑程度，是传统图拉普拉斯增强方法的核心正则项。
        """,
    )
    add_equation(doc, "L = D - W,     z^T L z = Σ_(i,j∈E) w_ij (z_i - z_j)^2")

    doc.add_heading("2.5 基于AMR-GTV的噪声增强算法", level=2)
    doc.add_heading("2.5.1 基础图拉普拉斯增强模型", level=3)
    add_paragraph(
        doc,
        """
        基础图拉普拉斯方法通过二次型平滑约束修正含噪观测。其目标函数由数据保真项和图平滑项组成，闭式解可由线性方程组直接给出。该方法结构简单，但对离群点和物理边界不够鲁棒，容易在强噪声下过度传播异常观测。
        """,
    )
    add_equation(doc, "min_z ||z-y||_2^2 + λ z^T L z,     z = (I + λL)^(-1)y")
    doc.add_heading("2.5.2 RAGLE基础版鲁棒图增强模型", level=3)
    add_paragraph(
        doc,
        """
        RAGLE在图拉普拉斯框架中引入观测置信度和鲁棒边权。观测置信度用于降低疑似噪声点的数据保真强度，鲁棒边权用于削弱数值差异过大的相邻点之间的平滑作用。相比普通图拉普拉斯，该方法对异常点更稳健，但仍主要依赖单尺度邻域图和二次型平滑。
        """,
    )
    add_equation(doc, "min_z Σ_i c_i(z_i-y_i)^2 + λΣ_(i,j∈E) w_ij q_ij(z_i-z_j)^2")
    doc.add_heading("2.5.3 本文AMR-GTV改进算法", level=3)
    add_paragraph(
        doc,
        """
        为进一步适应未知异方差噪声和离群观测，本文提出AMR-GTV算法。该算法包含四个关键设计：其一，构建多尺度kNN图以同时描述局部细节和较大尺度连续性；其二，根据含噪观测估计点级噪声尺度并形成观测置信度c_i；其三，采用Charbonnier形式的图总变差惩罚替代纯二次平滑，以增强边缘保持能力；其四，引入自适应中值先验，低置信度观测点更多依赖局部鲁棒统计，高置信度点则保持较强数据保真。
        """,
    )
    add_equation(
        doc,
        "min_z 0.5Σ_i c_i(z_i-y_i)^2 + λΣ_sΣ_(i,j∈E_s)w_ij^sρ_δ(z_i-z_j) + Σ_i η_i(z_i-z_med,i)^2",
    )
    add_equation(doc, "η_i = η(1-c_i),     ρ_δ(r)=sqrt(r^2+δ^2)")
    add_paragraph(
        doc,
        """
        上式中，η_i为点自适应中值先验权重。当c_i较小时，该点被认为可靠性较低，中值先验权重增大；当c_i接近1时，该点主要由数据保真项控制。实现时采用迭代重加权最小二乘方法求解，每次迭代根据当前估计更新图边权，并求解线性系统得到新的增强观测值。
        """,
    )
    add_table(
        doc,
        "表2-1 AMR-GTV噪声增强算法流程",
        ["步骤", "操作"],
        [
            ["1", "输入稀疏观测坐标、含噪观测值和邻域尺度集合"],
            ["2", "构建多尺度kNN图，并计算空间距离权重"],
            ["3", "基于局部中值残差估计点级噪声尺度和观测置信度"],
            ["4", "根据置信度构造数据保真矩阵和自适应中值先验"],
            ["5", "采用IRLS更新图总变差边权并求解线性系统"],
            ["6", "输出增强后的稀疏观测值，用于后续流场重建"],
        ],
    )
    doc.add_heading("2.5.4 算法复杂度分析", level=3)
    add_paragraph(
        doc,
        """
        设稀疏观测点数为N，邻域规模为k，多尺度图数量为S。kNN图构建需要计算点间距离，直接实现的复杂度为O(N^2)，但本章观测点规模较小，该开销可接受。每次IRLS迭代需要组装图拉普拉斯并求解N维线性系统。由于增强过程只发生在稀疏观测点上，而不是完整网格上，因此其计算代价远低于对全场进行高维优化。
        """,
    )

    doc.add_heading("2.6 噪声增强算法实验设计", level=2)
    doc.add_heading("2.6.1 实验数据与采样设置", level=3)
    add_paragraph(
        doc,
        """
        本章实验使用渤海区域流场数据中的temperature字段作为标量物理量。主压力实验选择两个时间片、三个采样率、四个噪声强度和两个随机种子，共形成48组主要实验条件。所有采样点均要求位于有效海域掩膜内。实验输出包括增强后的稀疏观测、误差指标表和论文图像。
        """,
    )
    add_table(
        doc,
        "表2-2 代表性压力实验参数设置",
        ["参数", "设置"],
        [
            ["数据文件", "data/current/bohai.npz"],
            ["物理量字段", "temperature"],
            ["采样率", "0.005、0.01、0.03"],
            ["噪声强度", "0.40、0.60、0.80、1.00"],
            ["噪声模型", "异方差离群噪声为主，另设噪声模型对比"],
            ["离群点比例", "主实验0.10，额外考察0、0.05、0.10、0.15"],
            ["随机种子", "3407、3408"],
            ["AMR-GTV尺度", "k=3、5、8"],
        ],
    )
    doc.add_heading("2.6.2 对比方法设置", level=3)
    add_paragraph(
        doc,
        """
        本章选择六种代表性方法作为对比对象。它们分别对应无处理基线、传统局部均值、鲁棒局部统计、经典图平滑、基础鲁棒图增强以及本文改进算法。这样的设计能够覆盖本章关注的主要技术路线，并使后续图表分析围绕清晰的算法层次展开。
        """,
    )
    add_table(
        doc,
        "表2-3 代表性对比算法说明",
        ["方法", "类型", "作用"],
        [
            ["Noisy", "无增强基线", "衡量噪声直接输入时的误差水平"],
            ["KNN-Mean", "传统局部平均", "反映简单邻域统计平滑效果"],
            ["Median", "鲁棒局部统计", "抑制局部离群点影响"],
            ["Graph-Laplacian", "经典图平滑", "检验二次型图正则的作用"],
            ["RAGLE", "鲁棒图拉普拉斯", "作为本文方法的基础版对比"],
            ["AMR-GTV", "本文方法", "结合点级置信度、多尺度图总变差和自适应中值先验"],
        ],
    )
    doc.add_heading("2.6.3 评价指标", level=3)
    add_paragraph(
        doc,
        """
        本章主文采用均方根误差RMSE、相对含噪观测的提升率以及运行时间评价增强效果。RMSE对局部大误差更敏感，适合衡量离群噪声残留情况；提升率用于展示增强算法相对于不处理基线的实际收益；运行时间用于说明不同方法的工程代价。实验明细文件中保留MAE和PSNR作为辅助记录，但正文图表与结论均围绕上述三个主指标展开。
        """,
    )
    add_equation(doc, "RMSE = sqrt(1/N Σ_i (z_i-y_i^*)^2)")
    add_equation(doc, "Gain = (RMSE_noisy - RMSE_method) / RMSE_noisy × 100%")

    doc.add_heading("2.7 噪声增强实验结果与分析", level=2)
    doc.add_heading("2.7.1 总体性能分析", level=3)
    add_paragraph(
        doc,
        f"""
        本节按照总体性能、噪声强度、采样率、鲁棒性、参数敏感性和运行效率展开分析。由总体结果可知，AMR-GTV取得最低平均RMSE，为{fmt(amr['rmse_mean'])}，相对于含噪观测的平均RMSE {fmt(noisy['rmse_mean'])} 降低约{pct((noisy['rmse_mean']-amr['rmse_mean'])/noisy['rmse_mean']*100)}%。相对于中值滤波、图拉普拉斯和RAGLE基础版，AMR-GTV的平均RMSE分别降低约{pct((median['rmse_mean']-amr['rmse_mean'])/median['rmse_mean']*100)}%、{pct((graph['rmse_mean']-amr['rmse_mean'])/graph['rmse_mean']*100)}%和{pct((ragle['rmse_mean']-amr['rmse_mean'])/ragle['rmse_mean']*100)}%。这说明本文方法并非只在单一场景中有效，而是在多噪声强度、多采样率和多随机种子组合下具有稳定优势。
        """,
    )
    add_table(
        doc,
        "表2-4 代表性算法压力实验总体结果",
        ["方法", "平均RMSE", "RMSE标准差", "平均提升率/%", "AMR相对降低率/%"],
        rows_ranking(ranking),
    )
    add_picture(
        doc,
        "fig_2_5_method_ranking.png",
        "图2-5 代表性方法在困难观测场景下的总体性能",
        "图2-5按平均RMSE对六种代表算法进行排序。可以看到，AMR-GTV处于误差最低的位置，且相对于RAGLE基础版和含噪观测具有明显差距。中值滤波和图拉普拉斯在部分场景中有效，但总体上仍弱于本文方法，说明单纯局部鲁棒统计或二次型图平滑都不足以充分处理异方差离群噪声。",
    )
    add_picture(
        doc,
        "fig_2_6_amr_gap_bar.png",
        "图2-6 AMR-GTV相对代表性算法的性能差距",
        "图2-6进一步给出AMR-GTV相对各基线方法的RMSE降低率。相对于RAGLE基础版、中值滤波和图拉普拉斯，AMR-GTV在总体均值上保持更低RMSE，说明当前参数设置下多尺度图约束、点级置信度和中值先验的组合具有较好的稳健性。该结果用于支撑主对比结论，而不单独展开组件拆分分析。",
    )
    add_picture(
        doc,
        "fig_2_7_error_distribution.png",
        "图2-7 代表性场景下的误差分布",
        "图2-7通过误差分布展示各方法对局部大误差的抑制能力。AMR-GTV不仅降低平均误差，还压缩了高误差尾部。该特性对于后续流场重建尤其重要，因为少量高误差观测点可能在连续场拟合中形成局部伪结构。",
    )
    add_picture(
        doc,
        "fig_2_8_spatial_error_maps.png",
        "图2-8 代表性场景下的空间误差分布",
        "图2-8从空间位置上比较误差分布。含噪观测和RAGLE仍存在明显高误差团簇，图拉普拉斯可降低部分随机扰动但边缘保持能力有限；AMR-GTV的高误差区域更少且分布更均匀，更适合作为后续重建模型输入。",
    )

    doc.add_heading("2.7.2 噪声强度与离群点影响", level=3)
    add_paragraph(
        doc,
        """
        噪声强度直接决定观测值偏离真实值的程度。表2-5和图2-9给出了不同噪声强度下的RMSE变化。低噪声条件下，各类增强算法之间差距有限；随着噪声强度增加，含噪观测和RAGLE基础版误差快速上升，而AMR-GTV曲线增长更缓，说明点级置信度和自适应中值先验能够在强噪声条件下有效削弱异常观测影响。
        """,
    )
    add_table(
        doc,
        "表2-5 不同噪声强度下代表性方法RMSE对比",
        ["噪声强度"] + [METHOD_NAMES[m] for m in METHOD_ORDER],
        pivot_rows(metrics["noise"], "noise_ratio", "rmse_mean"),
    )
    add_picture(
        doc,
        "fig_2_9_noise_ratio_rmse.png",
        "图2-9 不同强噪声水平下的RMSE变化",
        "图2-9表明，AMR-GTV在0.60和0.80等中高噪声条件下优势更加明显。在噪声强度为1.00时，中值滤波具有很强竞争力，但AMR-GTV仍保持相近甚至更稳定的总体表现，说明本文方法在强鲁棒局部统计基础上进一步利用了图结构信息。",
    )
    add_picture(
        doc,
        "fig_2_10_outlier_fraction_rmse.png",
        "图2-10 离群点比例变化下的误差曲线",
        "图2-10专门考察离群点比例变化。无离群点时，简单平滑或RAGLE可能已经足够；当离群点比例升高后，AMR-GTV逐渐体现优势。这一结果说明本文方法的主要适用场景是存在复杂观测误差的实际传感器数据，而不是理想轻噪声数据。",
    )

    doc.add_heading("2.7.3 不同采样率下的增强效果", level=3)
    add_paragraph(
        doc,
        """
        采样率影响邻域图的可靠性。采样率过低时，观测点之间距离较大，kNN图可能连接物理差异较明显的点，所有基于邻域的算法都会受到限制；采样率提高后，邻域关系更能反映真实局部连续性，图约束和局部统计都更稳定。
        """,
    )
    add_table(
        doc,
        "表2-6 不同采样率下代表性方法平均提升率/%",
        ["采样率"] + [METHOD_NAMES[m] for m in METHOD_ORDER],
        pivot_rows(metrics["sampling"], "sampling_rate", "gain_mean"),
    )
    add_picture(
        doc,
        "fig_2_11_sampling_gain.png",
        "图2-11 不同采样率下的增强收益",
        "图2-11显示，采样率从0.005提高至0.03后，AMR-GTV的收益更加明显。这说明多尺度图结构需要一定观测密度支撑；当邻域信息足够时，本文方法能够更好地区分真实空间变化与噪声扰动。",
    )
    add_picture(
        doc,
        "fig_2_12_amr_gain_heatmap.png",
        "图2-12 AMR-GTV相对图拉普拉斯的RMSE降低率",
        "图2-12以采样率和噪声强度为二维条件，展示AMR-GTV相对图拉普拉斯的提升。热力图中较高提升区域主要集中在中高噪声且邻域信息较充分的条件下，说明本文改进对普通二次型图平滑的优势具有明确适用边界。",
    )

    doc.add_heading("2.7.4 参数敏感性与稳定性分析", level=3)
    add_paragraph(
        doc,
        """
        为分析AMR-GTV内部参数对结果的影响，本文进一步考察图正则强度和自适应中值先验强度。参数敏感性分析的目的不是在测试集上寻找最优偶然参数，而是观察默认参数附近的性能变化趋势，判断所选参数是否处于相对稳定、低误差的区域。若中值先验过弱，离群扰动可能保留较多；若先验过强，则可能过度削弱真实局部变化。
        """,
    )
    add_picture(
        doc,
        "fig_2_13_parameter_sensitivity.png",
        "图2-13 AMR-GTV关键参数敏感性",
        "图2-13显示，较弱或过强的中值先验都会影响RMSE表现；适度中值先验结合图总变差约束后误差较低。红色星标为本文默认参数，处于低误差区域，说明参数选择具有合理性和一定稳定性。该结果用于说明参数设置口径，而不作为单独组件结论。",
    )
    add_picture(
        doc,
        "fig_2_14_casewise_reduction.png",
        "图2-14 AMR-GTV相对基线的逐条件提升分布",
        "图2-14统计每个实验条件下AMR-GTV相对各基线的RMSE降低率。多数条件下，AMR-GTV相对KNN均值、图拉普拉斯和RAGLE均取得正收益；相对中值滤波的提升幅度较小但整体仍为正，说明本文方法是在强鲁棒局部统计基础上的进一步改进。",
    )
    doc.add_heading("2.7.5 运行时间与工程代价分析", level=3)
    add_paragraph(
        doc,
        """
        除增强精度外，预处理算法的运行代价也会影响其工程可用性。本章所研究的增强过程仅在稀疏观测点上执行，不涉及完整网格的高维优化；因此即使AMR-GTV需要多尺度图构建和迭代求解，其计算规模仍由观测点数量决定。为考察该特性，本文统计不同观测规模下六种代表算法的平均运行时间。
        """,
    )
    add_picture(
        doc,
        "fig_2_15_runtime_vs_points.png",
        "图2-15 不同观测规模下的运行时间对比",
        "图2-15给出运行时间随观测点数量变化的趋势。AMR-GTV由于需要多尺度图构建和IRLS求解，时间开销高于KNN均值和中值滤波，但本章只在稀疏观测点上进行预处理，运行时间仍处于可接受范围。对于离线流场重建任务，该计算代价通常可以由更低的输入误差抵消。",
    )

    doc.add_heading("2.8 本章小结", level=2)
    add_paragraph(
        doc,
        """
        本章围绕含噪稀疏观测条件下的流场预处理问题，建立了从观测点表示、噪声建模、邻域图构建到噪声增强算法的完整方法。针对实际观测中噪声强度未知且各点不一致的问题，本文提出AMR-GTV算法，通过点级噪声估计构造观测置信度，并结合多尺度图总变差和自适应中值先验实现鲁棒增强。
        实验部分将主对比方法收敛为六种代表性算法，避免了大量相近算法造成的结论分散。结果表明，在异方差离群噪声、强噪声和不同采样率组合的压力实验中，AMR-GTV取得最低平均RMSE，相对中值滤波、图拉普拉斯和RAGLE基础版均有稳定提升。参数敏感性和逐条件提升分布进一步说明，本文采用的默认参数设置在多种压力条件下具有较好的稳定性。
        本章输出的增强稀疏观测仍保持原始观测点坐标不变，仅修正观测点物理量数值，因此可直接作为第三章连续流场隐式神经表示模型的输入。后续章节将在此基础上研究如何由增强后的稀疏观测重建完整连续流场。
        """,
    )

    doc.add_heading("本章参考文献建议", level=2)
    refs = [
        "[1] Shuman D I, Narang S K, Frossard P, Ortega A, Vandergheynst P. The emerging field of signal processing on graphs[J]. IEEE Signal Processing Magazine, 2013, 30(3): 83-98.",
        "[2] Wang Y X, Sharpnack J, Smola A J, Tibshirani R J. Trend Filtering on Graphs[J]. Journal of Machine Learning Research, 2016, 17(105): 1-41.",
        "[3] Kim S J, Koh K, Boyd S, Gorinevsky D. l1 Trend Filtering[J]. SIAM Review, 2009, 51(2): 339-360.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(ref)
        set_font(run)

    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    return DOCX


if __name__ == "__main__":
    print(build_doc())
