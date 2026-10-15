from __future__ import annotations

import copy
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "第二章_含噪稀疏观测建模与噪声增强算法_论文整合版.docx"
OUTPUT = ROOT / "formatted_outputs" / "第二章_含噪稀疏观测建模与噪声增强算法_按哈工大格式公式图片修正版.docx"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_MML = "http://www.w3.org/1998/Math/MathML"


EQUATIONS = {
    "O_t = {(x_i, y_i) | i = 1, 2, ..., N, x_i ∈ Ω_valid}": (
        r"""
        <mrow>
          <msub><mi>O</mi><mi>t</mi></msub><mo>=</mo>
          <mo>{</mo><mo>(</mo><msub><mi>x</mi><mi>i</mi></msub><mo>,</mo>
          <msub><mi>y</mi><mi>i</mi></msub><mo>)</mo><mo>|</mo>
          <mi>i</mi><mo>=</mo><mn>1</mn><mo>,</mo><mn>2</mn><mo>,</mo><mo>…</mo><mo>,</mo><mi>N</mi><mo>,</mo>
          <msub><mi>x</mi><mi>i</mi></msub><mo>∈</mo><msub><mi>Ω</mi><mtext>valid</mtext></msub><mo>}</mo>
        </mrow>
        """
    ),
    "y_i = y_i^* + ε_i,     ε_i ~ N(0, σ_i^2)": (
        r"""
        <mrow>
          <msub><mi>y</mi><mi>i</mi></msub><mo>=</mo>
          <msubsup><mi>y</mi><mi>i</mi><mo>*</mo></msubsup><mo>+</mo>
          <msub><mi>ε</mi><mi>i</mi></msub><mo>,</mo>
          <mspace width="1em"/>
          <msub><mi>ε</mi><mi>i</mi></msub><mo>∼</mo>
          <mi>N</mi><mo>(</mo><mn>0</mn><mo>,</mo><msubsup><mi>σ</mi><mi>i</mi><mn>2</mn></msubsup><mo>)</mo>
        </mrow>
        """
    ),
    "w_ij = exp(-d_ij^2 / (2τ^2)),     (i,j) ∈ E": (
        r"""
        <mrow>
          <msub><mi>w</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>=</mo>
          <mi mathvariant="normal">exp</mi><mo>(</mo>
          <mfrac><mrow><mo>-</mo><msubsup><mi>d</mi><mrow><mi>i</mi><mi>j</mi></mrow><mn>2</mn></msubsup></mrow>
          <mrow><mn>2</mn><msup><mi>τ</mi><mn>2</mn></msup></mrow></mfrac>
          <mo>)</mo><mo>,</mo><mspace width="1em"/>
          <mo>(</mo><mi>i</mi><mo>,</mo><mi>j</mi><mo>)</mo><mo>∈</mo><mi>E</mi>
        </mrow>
        """
    ),
    "L = D - W,     z^T L z = Σ_(i,j∈E) w_ij (z_i - z_j)^2": (
        r"""
        <mrow>
          <mi>L</mi><mo>=</mo><mi>D</mi><mo>-</mo><mi>W</mi><mo>,</mo><mspace width="1em"/>
          <msup><mi>z</mi><mi>T</mi></msup><mi>L</mi><mi>z</mi><mo>=</mo>
          <munderover><mo>∑</mo><mrow><mo>(</mo><mi>i</mi><mo>,</mo><mi>j</mi><mo>)</mo><mo>∈</mo><mi>E</mi></mrow><mrow/></munderover>
          <msub><mi>w</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub>
          <msup><mrow><mo>(</mo><msub><mi>z</mi><mi>i</mi></msub><mo>-</mo><msub><mi>z</mi><mi>j</mi></msub><mo>)</mo></mrow><mn>2</mn></msup>
        </mrow>
        """
    ),
    "min_z ||z-y||_2^2 + λ z^T L z,     z = (I + λL)^(-1)y": (
        r"""
        <mrow>
          <msub><mi mathvariant="normal">min</mi><mi>z</mi></msub>
          <msubsup><mrow><mo>‖</mo><mi>z</mi><mo>-</mo><mi>y</mi><mo>‖</mo></mrow><mn>2</mn><mn>2</mn></msubsup>
          <mo>+</mo><mi>λ</mi><msup><mi>z</mi><mi>T</mi></msup><mi>L</mi><mi>z</mi><mo>,</mo><mspace width="1em"/>
          <mi>z</mi><mo>=</mo>
          <msup><mrow><mo>(</mo><mi>I</mi><mo>+</mo><mi>λ</mi><mi>L</mi><mo>)</mo></mrow><mrow><mo>-</mo><mn>1</mn></mrow></msup><mi>y</mi>
        </mrow>
        """
    ),
    "min_z Σ_i c_i(z_i-y_i)^2 + λΣ_(i,j∈E) w_ij q_ij(z_i-z_j)^2": (
        r"""
        <mrow>
          <msub><mi mathvariant="normal">min</mi><mi>z</mi></msub>
          <msub><mo>∑</mo><mi>i</mi></msub>
          <msub><mi>c</mi><mi>i</mi></msub>
          <msup><mrow><mo>(</mo><msub><mi>z</mi><mi>i</mi></msub><mo>-</mo><msub><mi>y</mi><mi>i</mi></msub><mo>)</mo></mrow><mn>2</mn></msup>
          <mo>+</mo><mi>λ</mi>
          <munderover><mo>∑</mo><mrow><mo>(</mo><mi>i</mi><mo>,</mo><mi>j</mi><mo>)</mo><mo>∈</mo><mi>E</mi></mrow><mrow/></munderover>
          <msub><mi>w</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub>
          <msub><mi>q</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub>
          <msup><mrow><mo>(</mo><msub><mi>z</mi><mi>i</mi></msub><mo>-</mo><msub><mi>z</mi><mi>j</mi></msub><mo>)</mo></mrow><mn>2</mn></msup>
        </mrow>
        """
    ),
    "min_z 0.5Σ_i c_i(z_i-y_i)^2 + λΣ_sΣ_(i,j∈E_s)w_ij^sρ_δ(z_i-z_j) + Σ_i η_i(z_i-z_med,i)^2": (
        r"""
        <mrow>
          <msub><mi mathvariant="normal">min</mi><mi>z</mi></msub>
          <mn>0.5</mn><msub><mo>∑</mo><mi>i</mi></msub>
          <msub><mi>c</mi><mi>i</mi></msub>
          <msup><mrow><mo>(</mo><msub><mi>z</mi><mi>i</mi></msub><mo>-</mo><msub><mi>y</mi><mi>i</mi></msub><mo>)</mo></mrow><mn>2</mn></msup>
          <mo>+</mo><mi>λ</mi><msub><mo>∑</mo><mi>s</mi></msub>
          <munderover><mo>∑</mo><mrow><mo>(</mo><mi>i</mi><mo>,</mo><mi>j</mi><mo>)</mo><mo>∈</mo><msub><mi>E</mi><mi>s</mi></msub></mrow><mrow/></munderover>
          <msubsup><mi>w</mi><mrow><mi>i</mi><mi>j</mi></mrow><mi>s</mi></msubsup>
          <msub><mi>ρ</mi><mi>δ</mi></msub><mo>(</mo><msub><mi>z</mi><mi>i</mi></msub><mo>-</mo><msub><mi>z</mi><mi>j</mi></msub><mo>)</mo>
          <mo>+</mo><msub><mo>∑</mo><mi>i</mi></msub>
          <msub><mi>η</mi><mi>i</mi></msub>
          <msup><mrow><mo>(</mo><msub><mi>z</mi><mi>i</mi></msub><mo>-</mo><msub><mi>z</mi><mrow><mtext>med</mtext><mo>,</mo><mi>i</mi></mrow></msub><mo>)</mo></mrow><mn>2</mn></msup>
        </mrow>
        """
    ),
    "η_i = η(1-c_i),     ρ_δ(r)=sqrt(r^2+δ^2)": (
        r"""
        <mrow>
          <msub><mi>η</mi><mi>i</mi></msub><mo>=</mo><mi>η</mi><mo>(</mo><mn>1</mn><mo>-</mo><msub><mi>c</mi><mi>i</mi></msub><mo>)</mo><mo>,</mo><mspace width="1em"/>
          <msub><mi>ρ</mi><mi>δ</mi></msub><mo>(</mo><mi>r</mi><mo>)</mo><mo>=</mo>
          <msqrt><mrow><msup><mi>r</mi><mn>2</mn></msup><mo>+</mo><msup><mi>δ</mi><mn>2</mn></msup></mrow></msqrt>
        </mrow>
        """
    ),
    "RMSE = sqrt(1/N Σ_i (z_i-y_i^*)^2)": (
        r"""
        <mrow>
          <mi mathvariant="normal">RMSE</mi><mo>=</mo>
          <msqrt><mrow>
            <mfrac><mn>1</mn><mi>N</mi></mfrac>
            <msub><mo>∑</mo><mi>i</mi></msub>
            <msup><mrow><mo>(</mo><msub><mi>z</mi><mi>i</mi></msub><mo>-</mo><msubsup><mi>y</mi><mi>i</mi><mo>*</mo></msubsup><mo>)</mo></mrow><mn>2</mn></msup>
          </mrow></msqrt>
        </mrow>
        """
    ),
    "Gain = (RMSE_noisy - RMSE_method) / RMSE_noisy × 100%": (
        r"""
        <mrow>
          <mi mathvariant="normal">Gain</mi><mo>=</mo>
          <mfrac>
            <mrow><msub><mi mathvariant="normal">RMSE</mi><mtext>noisy</mtext></msub><mo>-</mo><msub><mi mathvariant="normal">RMSE</mi><mtext>method</mtext></msub></mrow>
            <msub><mi mathvariant="normal">RMSE</mi><mtext>noisy</mtext></msub>
          </mfrac>
          <mo>×</mo><mn>100</mn><mo>%</mo>
        </mrow>
        """
    ),
}


INLINE_MATH = {
    "AMR-GTV": None,
    "IRLS": None,
    "RMSE": r"<mi mathvariant='normal'>RMSE</mi>",
    "MAE": r"<mi mathvariant='normal'>MAE</mi>",
    "Gain": r"<mi mathvariant='normal'>Gain</mi>",
    "N=round(rHW)": r"<mrow><mi>N</mi><mo>=</mo><mi mathvariant='normal'>round</mi><mo>(</mo><mi>r</mi><mi>H</mi><mi>W</mi><mo>)</mo></mrow>",
    "u(x)": r"<mrow><mi>u</mi><mo>(</mo><mi>x</mi><mo>)</mo></mrow>",
    "O_t": r"<msub><mi>O</mi><mi>t</mi></msub>",
    "{(x_i,y_i)}": r"<mrow><mo>{</mo><mo>(</mo><msub><mi>x</mi><mi>i</mi></msub><mo>,</mo><msub><mi>y</mi><mi>i</mi></msub><mo>)</mo><mo>}</mo></mrow>",
    "y_i^*": r"<msubsup><mi>y</mi><mi>i</mi><mo>*</mo></msubsup>",
    "y_i": r"<msub><mi>y</mi><mi>i</mi></msub>",
    "x_i": r"<msub><mi>x</mi><mi>i</mi></msub>",
    "x_j": r"<msub><mi>x</mi><mi>j</mi></msub>",
    "z_i": r"<msub><mi>z</mi><mi>i</mi></msub>",
    "z_j": r"<msub><mi>z</mi><mi>j</mi></msub>",
    "z_med,i": r"<msub><mi>z</mi><mrow><mtext>med</mtext><mo>,</mo><mi>i</mi></mrow></msub>",
    "w_ij": r"<msub><mi>w</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub>",
    "d_ij": r"<msub><mi>d</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub>",
    "q_ij": r"<msub><mi>q</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub>",
    "c_i": r"<msub><mi>c</mi><mi>i</mi></msub>",
    "η_i": r"<msub><mi>η</mi><mi>i</mi></msub>",
    "ρ_δ": r"<msub><mi>ρ</mi><mi>δ</mi></msub>",
    "E_s": r"<msub><mi>E</mi><mi>s</mi></msub>",
    "σ_i": r"<msub><mi>σ</mi><mi>i</mi></msub>",
    "ε_i": r"<msub><mi>ε</mi><mi>i</mi></msub>",
    "λ": r"<mi>λ</mi>",
    "τ": r"<mi>τ</mi>",
    "δ": r"<mi>δ</mi>",
    "Ω": r"<mi>Ω</mi>",
    "N": r"<mi>N</mi>",
    "k": r"<mi>k</mi>",
    "r": r"<mi>r</mi>",
    "H": r"<mi>H</mi>",
    "W": r"<mi>W</mi>",
    "L": r"<mi>L</mi>",
    "D": r"<mi>D</mi>",
    "E": r"<mi>E</mi>",
    "z": r"<mi>z</mi>",
    "y": r"<mi>y</mi>",
    "x": r"<mi>x</mi>",
}


def mathml(body: str) -> bytes:
    return f'<math xmlns="{NS_MML}">{body}</math>'.encode("utf-8")


class MathConverter:
    def __init__(self) -> None:
        self.transform = etree.XSLT(etree.parse(str(MML2OMML)))

    def omml(self, body: str, display: bool = False):
        result = self.transform(etree.fromstring(mathml(body)))
        root = etree.fromstring(etree.tostring(result))
        if root.tag.endswith("oMathPara"):
            if display:
                return root
            found = root.find(f".//{{{NS_M}}}oMath")
            return found if found is not None else root
        if root.tag.endswith("oMath"):
            return root
        found = root.find(f".//{{{NS_M}}}oMath")
        if found is not None:
            return found
        raise ValueError(f"Could not convert MathML body: {body[:120]}")


def set_rfonts(run, east_asia: str = "宋体", ascii_font: str = "Times New Roman") -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:ascii"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)


def set_paragraph_font(paragraph, size: float = 12, east_asia: str = "宋体", bold: bool = False) -> None:
    for run in paragraph.runs:
        set_rfonts(run, east_asia=east_asia)
        run.font.size = Pt(size)
        run.bold = bold


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def keep_lines(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepLines")) is None:
        p_pr.append(OxmlElement("w:keepLines"))


def add_math_to_paragraph(paragraph, converter: MathConverter, body: str) -> None:
    paragraph._p.append(copy.deepcopy(converter.omml(body)))


def add_text_run(paragraph, text: str, size: float = 12, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    set_rfonts(run)
    run.font.size = Pt(size)
    run.bold = bold


def format_body_paragraph(paragraph) -> None:
    fmt = paragraph.paragraph_format
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0.85)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(20)
    set_paragraph_font(paragraph, 12)


def apply_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.top_margin = Cm(2.85)
    section.bottom_margin = Cm(2.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.85)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(20)

    specs = {
        "Heading 1": (18, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 2": (15, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 3": (14, WD_ALIGN_PARAGRAPH.LEFT),
    }
    for name, (size, align) in specs.items():
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = False
        style.paragraph_format.alignment = align
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True


def normalize_heading_text(text: str) -> str:
    text = re.sub(r"^(第\d+章)\s*", r"\1  ", text)
    text = re.sub(r"^(\d+\.\d+(?:\.\d+)?)\s*", r"\1  ", text)
    return text


def format_headings(doc: Document) -> None:
    for i, p in enumerate(doc.paragraphs):
        if p.style.name == "Heading 1":
            p.text = normalize_heading_text(p.text.strip())
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            keep_with_next(p)
            set_paragraph_font(p, 18, "黑体", False)
            if i == 0:
                p.paragraph_format.page_break_before = True
        elif p.style.name == "Heading 2":
            p.text = normalize_heading_text(p.text.strip())
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            keep_with_next(p)
            set_paragraph_font(p, 15, "黑体", False)
        elif p.style.name == "Heading 3":
            p.text = normalize_heading_text(p.text.strip())
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            keep_with_next(p)
            set_paragraph_font(p, 14, "黑体", False)


def is_caption_text(text: str) -> bool:
    return bool(re.match(r"^[图表]2-\d+\s+", text.strip()))


def format_caption(paragraph) -> None:
    text = paragraph.text.strip()
    if text.startswith("表2-6"):
        text = text.replace("/%", "")
    text = re.sub(r"^([图表]2-\d+)\s+", r"\1  ", text)
    paragraph.text = text
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    keep_lines(paragraph)
    if text.startswith("表"):
        keep_with_next(paragraph)
    set_paragraph_font(paragraph, 10.5, "宋体", False)


def set_cell_shading(cell, fill: str = "FFFFFF") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, margin: int = 108) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin))
        node.set(qn("w:type"), "dxa")


def border_element(edge: str, val: str, size: int = 8) -> OxmlElement:
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), val)
    if val != "nil":
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
    return element


def set_table_three_line(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for existing in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    borders.append(border_element("top", "single", 12))
    borders.append(border_element("left", "nil"))
    borders.append(border_element("bottom", "single", 12))
    borders.append(border_element("right", "nil"))
    borders.append(border_element("insideH", "nil"))
    borders.append(border_element("insideV", "nil"))
    tbl_pr.append(borders)

    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for cell in table.rows[0].cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = tc_pr.find(qn("w:tcBorders"))
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                tc_pr.append(borders)
            for edge in ("top", "left", "right"):
                old = borders.find(qn(f"w:{edge}"))
                if old is not None:
                    borders.remove(old)
                borders.append(border_element(edge, "nil"))
            old = borders.find(qn("w:bottom"))
            if old is not None:
                borders.remove(old)
            borders.append(border_element("bottom", "single", 8))


def format_tables(doc: Document) -> None:
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        set_table_three_line(table)
        for r_idx, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_shading(cell)
                set_cell_margins(cell, 108)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    set_paragraph_font(p, 10.5, "宋体", r_idx == 0)


def split_inline_math(text: str):
    candidates = [k for k, v in INLINE_MATH.items() if v and k in text]
    candidates.sort(key=len, reverse=True)
    if not candidates:
        return [(False, text)]
    parts = []
    for candidate in candidates:
        escaped = re.escape(candidate)
        if re.fullmatch(r"[A-Za-z]+", candidate):
            escaped = rf"(?<![A-Za-z0-9_]){escaped}(?![A-Za-z0-9_])"
        elif len(candidate) == 1:
            escaped = rf"(?<![A-Za-z0-9_]){escaped}(?![A-Za-z0-9_])"
        parts.append(escaped)
    pattern = re.compile("|".join(parts))
    result = []
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            result.append((False, text[pos:match.start()]))
        result.append((True, match.group(0)))
        pos = match.end()
    if pos < len(text):
        result.append((False, text[pos:]))
    return result


def rebuild_inline_math(paragraph, converter: MathConverter) -> None:
    text = paragraph.text
    if not text or is_caption_text(text) or text.strip() in EQUATIONS:
        return
    if paragraph.style.name.startswith("Heading"):
        return
    pieces = split_inline_math(text)
    if all(not is_math for is_math, _ in pieces):
        format_body_paragraph(paragraph)
        return
    clear_paragraph(paragraph)
    for is_math, piece in pieces:
        if is_math:
            add_math_to_paragraph(paragraph, converter, INLINE_MATH[piece])
        else:
            add_text_run(paragraph, piece, 12)
    format_body_paragraph(paragraph)


def replace_equation_paragraph(paragraph, converter: MathConverter, body: str, number: int) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is not None:
        p_pr.remove(tabs)
    tabs = OxmlElement("w:tabs")
    center = OxmlElement("w:tab")
    center.set(qn("w:val"), "center")
    center.set(qn("w:pos"), "4252")
    right = OxmlElement("w:tab")
    right.set(qn("w:val"), "right")
    right.set(qn("w:pos"), "8504")
    tabs.append(center)
    tabs.append(right)
    p_pr.append(tabs)
    paragraph.add_run().add_tab()
    add_math_to_paragraph(paragraph, converter, body)
    paragraph.add_run().add_tab()
    add_text_run(paragraph, f"（2-{number}）", 12)
    keep_lines(paragraph)


def replace_equations(doc: Document, converter: MathConverter) -> None:
    eq_no = 1
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in EQUATIONS:
            replace_equation_paragraph(p, converter, EQUATIONS[text], eq_no)
            eq_no += 1


def update_formula_references(doc: Document) -> None:
    replacements = {
        "形成稀疏观测集合O_t={(x_i,y_i)}。": "形成稀疏观测集合O_t={(x_i,y_i)}，其数学定义见式（2-1）。",
        "本章实验同时考虑同方差高斯噪声、异方差高斯噪声和异方差离群噪声。": "本章实验同时考虑同方差高斯噪声、异方差高斯噪声和异方差离群噪声，加性噪声模型见式（2-2）。",
        "距离越近，空间权重越大；距离越远，空间权重越小。": "距离越近，空间权重越大；距离越远，空间权重越小，边权定义见式（2-3）。",
        "是传统图拉普拉斯增强方法的核心正则项。": "是传统图拉普拉斯增强方法的核心正则项，如式（2-4）所示。",
        "闭式解可由线性方程组直接给出。": "闭式解可由线性方程组直接给出，目标函数与闭式解见式（2-5）。",
        "但仍主要依赖单尺度邻域图和二次型平滑。": "但仍主要依赖单尺度邻域图和二次型平滑，其优化形式见式（2-6）。",
        "高置信度点则保持较强数据保真。": "高置信度点则保持较强数据保真。AMR-GTV的目标函数见式（2-7）。",
        "上式中": "式（2-8）中",
        "上述目标函数": "式（2-7）所示目标函数",
        "本章采用均方根误差RMSE、平均绝对误差MAE和相对含噪观测提升率Gain评价增强效果。": (
            "本章采用均方根误差RMSE、平均绝对误差MAE和相对含噪观测提升率Gain评价增强效果，"
            "其中RMSE和Gain分别见式（2-9）和式（2-10）。"
        ),
    }
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") or is_caption_text(p.text):
            continue
        text = p.text
        new_text = text
        for old, new in replacements.items():
            new_text = new_text.replace(old, new)
        if new_text != text:
            p.text = new_text
            format_body_paragraph(p)


def mark_picture_paragraphs(doc: Document) -> None:
    for p in doc.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            # Picture paragraphs must override Normal's exact line spacing,
            # otherwise Word clips inline images to one text line.
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.line_spacing = 1.0
            keep_with_next(p)
            keep_lines(p)


def format_paragraphs(doc: Document, converter: MathConverter) -> None:
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") or p.text.strip() in EQUATIONS:
            continue
        if is_caption_text(p.text):
            format_caption(p)
        elif p._p.xpath(".//w:drawing"):
            continue
        elif p.text.strip():
            rebuild_inline_math(p, converter)


def final_audit(path: Path) -> dict[str, int]:
    with zipfile.ZipFile(path) as zf:
        document_xml = zf.read("word/document.xml")
        embeddings = [name for name in zf.namelist() if name.startswith("word/embeddings/")]
    return {
        "omath": document_xml.count(b"<m:oMath"),
        "omath_para": document_xml.count(b"<m:oMathPara"),
        "embeddings": len(embeddings),
        "left_borders": document_xml.count(b"<w:left"),
        "inside_v_borders": document_xml.count(b"<w:insideV"),
    }


def main() -> None:
    converter = MathConverter()
    doc = Document(INPUT)
    apply_styles(doc)
    format_headings(doc)
    update_formula_references(doc)
    mark_picture_paragraphs(doc)
    format_paragraphs(doc, converter)
    replace_equations(doc, converter)
    format_tables(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    audit = final_audit(OUTPUT)
    print(OUTPUT)
    print(audit)


if __name__ == "__main__":
    main()
